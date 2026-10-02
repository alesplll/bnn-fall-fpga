"""Extend the supplied PR1 report with the verified PR2 results."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "source_pr1.docx"
OUTPUT = ROOT / "Отчет_ПР1_ПР2_Рябец_Головин_Гришин.docx"
SCREENSHOTS = ROOT / "screenshots"


def replace_text(paragraph, value: str) -> None:
    """Keep the first run's formatting, including the existing title page."""
    if not paragraph.runs:
        paragraph.add_run(value)
        return
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def before(anchor, text: str, style: str):
    paragraph = anchor.insert_paragraph_before(style=style)
    paragraph.add_run(text)
    return paragraph


def toc_before(anchor, text: str, template):
    paragraph = before(anchor, text, "Report Contents")
    paragraph._p.remove(paragraph._p.pPr)
    paragraph._p.insert(0, deepcopy(template._p.pPr))
    return paragraph


def table_before(document, anchor, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "000000")
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index, title in enumerate(headers):
        table.columns[index].width = Cm(widths[index])
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.text = title
        cell.paragraphs[0].style = document.styles["Report Table Text"]
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Cm(widths[index])
            cells[index].text = value
            cells[index].paragraphs[0].style = document.styles["Report Table Text"]
    anchor._p.addprevious(table._tbl)
    return table


def evidence_figure_before(anchor, filename: str, caption: str) -> None:
    paragraph = anchor.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(SCREENSHOTS / filename), width=Cm(15.5))
    label = anchor.insert_paragraph_before(caption, style="Report Body")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(12)


def main() -> None:
    doc = Document(SOURCE)
    p = doc.paragraphs

    replace_text(p[8], "ОТЧЁТ ПО ПРАКТИЧЕСКИМ РАБОТАМ № 1–2")
    replace_text(p[11], "Тема: Системная и RTL-модели бинарного нейросетевого классификатора для детектирования падений")
    replace_text(
        p[40],
        "Цель практических работ № 1–2 — построить программную эталонную модель "
        "нейроускорителя, затем разработать её RTL-эквивалент на Verilog и "
        "сверить результаты симуляции с эталонным откликом. Модель "
        "классифицирует подготовленные признаки движения и выдаёт один бит решения.",
    )
    replace_text(
        p[41],
        "Решается задача детектирования падения человека по пяти признакам "
        "из данных акселерометра и гироскопа. Сеть 5→8→1 использует знаковые "
        "восьмибитные входы и веса в первом слое; выходной слой работает с "
        "бинарными активациями и весами через XNOR и подсчёт совпадений. "
        "Такой вычислительный тракт допускает целочисленную реализацию на ПЛИС.",
    )
    before(p[48], "– реализовать целочисленный классификатор 5→8→1 на Verilog с загрузкой весов из .mem-файлов;", "Report Body")
    before(p[48], "– выполнить симуляцию RTL-модели и проверить входы, промежуточные состояния и выходы по эталонным векторам.", "Report Body")

    replace_text(
        p[50],
        "Для детектирования падений разработана программная системная модель "
        "нейроускорителя и сформирован эталонный отклик — контрольные пары "
        "«вход — выход». По ним проверяется эквивалентность RTL-модели. "
        "Вычисления классификатора выполняются в целочисленной арифметике, "
        "поэтому результат можно воспроизвести бит в бит в Verilog без "
        "операций с плавающей точкой.",
    )
    replace_text(
        p[58],
        p[58].text.replace("Согласно постановке задачи проекта, для классификации используется 5 признаков.",
                           "Для классификации используются пять признаков."),
    )
    replace_text(
        p[77],
        "где H = 8 — число нейронов слоя A. Тождество проверено численно на "
        "2000 случайных парах ±1-векторов разной длины; во всех случаях "
        "получено полное совпадение. Аппаратная часть ограничена классификатором: "
        "слой B реализуется через XNOR, подсчёт единиц и пороговое сравнение, "
        "а слой A — через целочисленное умножение и накопление. Всего "
        "обучаемых параметров: 40 (W_A, int8) + 8 (b_A, int32) + "
        "8 (W_B, 1 бит) + 1 (b_B, int32) = 57 скалярных величин.",
    )
    replace_text(
        p[99],
        p[99].text.replace(" — для системы безопасности это стоит улучшать в первую очередь (например, решением по нескольким последовательным окнам), что выходит за рамки практики 1.",
                           ". Для системы безопасности снижение числа пропусков возможно, например, путём объединения решений по нескольким последовательным окнам."),
    )

    replace_text(p[100], "5.2 Эталонные данные для верификации RTL-модели")
    replace_text(
        p[101],
        "Для проверки RTL-модели подготовлен эталонный набор данных:",
    )
    replace_text(p[102], "– полный тестовый набор из 2331 образца: пять квантованных признаков, восьмибитный код слоя A, целочисленная сумма слоя B и итоговый бит;")
    replace_text(p[103], "– контрольный набор из 64 пар «40-битный вход — эталонный выход» для начальной симуляции: 32 падения и 32 обычные активности;")
    replace_text(p[104], "– параметры классификатора: 40 знаковых весов и 8 смещений слоя A, 8 бинарных весов и одно смещение слоя B;")
    replace_text(p[105], "– шестнадцатеричные файлы памяти для загрузки параметров и входных векторов в Verilog через $readmemh.")
    replace_text(p[106], "Контрольный набор проверен повторным декодированием входных слов в int8 и прогоном через целочисленную модель: все 64 ожидаемых выхода воспроизведены.")
    replace_text(p[107], "5.3 Границы аппаратной реализации")
    replace_text(p[121], "Ключевые алгоритмы программной модели представлены следующими фрагментами.")

    # The source document has a static contents page. Page numbers are filled
    # after rendering, while these entries preserve its paragraph style.
    replace_text(p[31], "2 ИЗВЛЕЧЕНИЕ ПРИЗНАКОВ И ФОРМИРОВАНИЕ ВЫБОРОК\t6")
    replace_text(p[32], "3 АРХИТЕКТУРА СЕТИ И КВАНТОВАНИЕ\t8")
    replace_text(p[33], "4 ОБУЧЕНИЕ И ПРОГРАММНАЯ РЕАЛИЗАЦИЯ\t10")
    replace_text(p[34], "5 РЕЗУЛЬТАТЫ И ЭТАЛОННЫЙ ОТКЛИК\t12")
    toc_anchor = p[35]
    toc_before(toc_anchor, "6 ПОСТАНОВКА ЗАДАЧИ ПРАКТИЧЕСКОЙ РАБОТЫ № 2\t15", p[34])
    toc_before(toc_anchor, "7 РЕАЛИЗАЦИЯ RTL-МОДЕЛИ\t17", p[34])
    toc_before(toc_anchor, "8 СИМУЛЯЦИЯ И СВЕРКА С ЭТАЛОНОМ\t18", p[34])
    replace_text(p[35], "ЗАКЛЮЧЕНИЕ\t21")
    replace_text(p[36], "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ\t22")
    replace_text(p[37], "ПРИЛОЖЕНИЕ А КЛЮЧЕВЫЕ ФРАГМЕНТЫ ПРОГРАММЫ\t23")
    toc_before(p[38], "ПРИЛОЖЕНИЕ Б ФРАГМЕНТЫ RTL-МОДЕЛИ\t28", p[34])

    anchor = p[109]  # Insert all PR2 material before the original conclusion.
    before(anchor, "6 ПОСТАНОВКА ЗАДАЧИ ПРАКТИЧЕСКОЙ РАБОТЫ № 2", "Report Heading 1")
    before(anchor, "6.1 Цель и границы RTL-модели", "Report Heading 2")
    before(anchor,
           "RTL-модель на Verilog реализует целочисленный классификатор, "
           "описанный в разделе 3. Пять рассчитанных и квантованных признаков "
           "подаются на вход; поиск пика, вычисление угла, jerk и дисперсии "
           "выполняются при программной подготовке данных. Совпадение выходов "
           "RTL-модели и программного эталона проверяется симуляцией.", "Report Body")
    before(anchor, "6.2 Интерфейс и формат данных", "Report Heading 2")
    before(anchor,
           "Входные признаки передаются одним 40-битным словом "
           "{x0, x1, x2, x3, x4}; каждый байт является знаковым числом int8 "
           "в дополнительном коде. По порядку передаются пиковое ускорение, "
           "угол наклона, пиковая угловая скорость, jerk и дисперсия после пика. "
           "Результат регистрируется на фронте тактового сигнала при valid_in=1. "
           "Сигнал valid_out показывает, что выходной бит относится к принятому "
           "входу. Сброс rst_n синхронный, активный низким уровнем.", "Report Body")
    before(anchor, "Таблица 6.1 — Интерфейс RTL-классификатора", "Report Table Caption")
    table_before(doc, anchor,
                 ["Сигнал", "Бит", "Назначение"],
                 [["clk", "1", "Тактовый сигнал"],
                  ["rst_n", "1", "Синхронный сброс, активный 0"],
                  ["valid_in", "1", "На входе подготовлено слово признаков"],
                  ["input_features", "40", "Пять знаковых байтов x0…x4"],
                  ["valid_out", "1", "Зарегистрированный ответ действителен"],
                  ["fall_detected", "1", "1 — падение, 0 — обычная активность"]],
                 [3.5, 2.7, 9.3])
    after_table = before(anchor,
           "Проект bnn_pr2 создан в Vivado 2022.2 для кристалла Xilinx "
           "Artix-7 xc7a100tcsg324-1. В его состав входят RTL-модуль, "
           "тестбенч и файлы памяти с параметрами и тестовыми данными "
           "(рисунки 6.1 и 6.2).", "Report Body")
    after_table.paragraph_format.space_before = Pt(6)
    evidence_figure_before(anchor, "vivado_project.jpg", "Рисунок 6.1 — Проект RTL-модели в Vivado 2022.2")
    evidence_figure_before(anchor, "vivado_rtl_code_interface.jpg", "Рисунок 6.2 — Интерфейс и загрузка параметров в коде Verilog")

    before(anchor, "7 РЕАЛИЗАЦИЯ RTL-МОДЕЛИ", "Report Heading 1")
    before(anchor, "7.1 Слой A: знаковое умножение и накопление", "Report Heading 2")
    before(anchor,
           "Восемь нейронов слоя A вычисляют z_h = b_A[h] + Σ_j x_j·W_A[j,h] "
           "для j=0…4. Веса W_A загружаются в порядке j·8+h. Операнды "
           "int8 явно трактуются как знаковые; произведение хранится в 16 битах "
           "и расширяется до 32 бит перед сложением со смещением int32. "
           "Промежуточный бит a_h равен 1 тогда и только тогда, когда z_h≥0. "
           "При нулевой сумме бит также равен 1, как и в программном эталоне.", "Report Body")
    before(anchor, "7.2 Слой B: XNOR и popcount", "Report Heading 2")
    before(anchor,
           "Восемь бинарных весов W_B сравниваются с битами a_h операцией "
           "XNOR. Четырёхбитный счётчик определяет число совпадений от 0 до 8. "
           "Итоговая сумма вычисляется как y = 2·popcount(XNOR(a,W_B)) − 8 + b_B; "
           "выход fall_detected равен 1 при y≥0. Эта формула бит в бит "
           "соответствует программной модели и не требует вещественной "
           "арифметики во время вывода.", "Report Body")
    before(anchor, "7.3 Регистры, сброс и загрузка параметров", "Report Heading 2")
    before(anchor,
           "Комбинационная часть рассчитывает оба слоя для текущего входа. "
           "На ближайшем фронте clk при valid_in=1 сохраняются выходной бит и "
           "valid_out=1; при отсутствии входа valid_out=0, а последний бит "
           "остаётся неизменным. Синхронный сброс обнуляет оба регистра. "
           "Веса и смещения загружаются из четырёх .mem-файлов через "
           "$readmemh: 40 весов W_A идут в порядке j·8+h, за ними отдельно "
           "хранятся 8 смещений b_A, 8 бинарных весов W_B и смещение b_B. "
           "Каждая строка файла памяти задаёт одно значение в hex-формате.", "Report Body")

    before(anchor, "8 СИМУЛЯЦИЯ И СВЕРКА С ЭТАЛОНОМ", "Report Heading 1")
    before(anchor, "8.1 Методика верификации", "Report Heading 2")
    before(anchor,
           "Самопроверяющийся тестбенч подаёт вход до фронта clk и "
           "сверяет зарегистрированный выход и valid_out с эталоном. На "
           "полном наборе он сравнивает также a_bits и сумму y, что "
           "локализует ошибки по слоям. Отдельно проверяются сброс и снятие "
           "valid_out. Полный набор содержит 2331 эталонную запись и "
           "преобразован в формат памяти для симуляции Verilog.", "Report Body")
    before(anchor,
           "Ожидаемый выход каждого вектора — решение программной модели, "
           "а не истинная метка SisFall. Симуляция проверяет точность "
           "аппаратного вычислительного тракта относительно целочисленного эталона.",
           "Report Body")
    before(anchor, "8.2 Результаты", "Report Heading 2")
    before(anchor, "Таблица 8.1 — Результаты поведенческой симуляции", "Report Table Caption")
    table_before(doc, anchor,
                 ["Набор", "Число", "Расхождения по выходу", "Расхождения внутри слоёв"],
                 [["Vivado XSim, полный набор", "2331", "0", "0 по a_bits и y"]],
                 [4.0, 2.1, 3.3, 6.1])
    after_table = before(anchor,
           "Например, для x=[14,16,27,2,0] модель и RTL дали "
           "a_bits=10100001, y=−5, fall_detected=0. Расхождений "
           "не обнаружено. Это подтверждает функциональную "
           "эквивалентность RTL целочисленному эталону на проверенных входах.", "Report Body")
    after_table.paragraph_format.space_before = Pt(6)
    before(anchor,
           "В Vivado 2022.2 создан проект bnn_pr2 для xc7a100tcsg324-1. "
           "Поведенческая симуляция XSim завершилась сообщением "
           "PASS: 2331 vectors (рисунок 8.1). Временные диаграммы показаны "
           "на рисунке 8.2. Синтез успешно завершился; получена логическая "
           "схема из 960 ячеек (рисунок 8.3). Согласно отчёту ресурсов "
           "(рисунок 8.4), "
           "использовано 629 Slice LUT, 2 Slice Register, 45 Bonded IOB "
           "и 1 BUFGCTRL. Оценка ресурсов относится только к синтезу; "
           "размещение, трассировка и анализ таймингов не выполнялись.", "Report Body")
    evidence_figure_before(anchor, "vivado_simulation_pass.jpg", "Рисунок 8.1 — Результат проверки 2331 эталонного вектора в XSim")
    evidence_figure_before(anchor, "vivado_waveform.jpg", "Рисунок 8.2 — Временные диаграммы входов и выходов классификатора")
    evidence_figure_before(anchor, "vivado_synthesized_schematic.jpg", "Рисунок 8.3 — Синтезированная логическая схема классификатора")
    evidence_figure_before(anchor, "vivado_utilization.jpg", "Рисунок 8.4 — Фрагмент схемы и ресурсы после синтеза")

    replace_text(
        p[110],
        "В практических работах № 1–2 создана программная и RTL-модель "
        "бинарного классификатора падений 5→8→1. Программная модель "
        "обучена на признаках SisFall и дала на тестовой выборке 97,0 % "
        "accuracy, 84,5 % precision и 93,6 % recall. RTL-модель повторяет "
        "целочисленный тракт: знаковые int8-умножения, 32-битное накопление, "
        "бинаризацию, XNOR и popcount. В Vivado 2022.2 создан проект "
        "для xc7a100tcsg324-1; симуляция XSim прошла 2331 вектор без "
        "расхождений по выходу и промежуточным состояниям обоих слоёв. "
        "Синтез завершился успешно и "
        "получена логическая схема. Таким образом, программная и RTL-модели "
        "согласованы по всем проверенным входам.",
    )

    source_anchor = p[119]
    before(source_anchor,
           "8. AMD. Vivado Design Suite User Guide: Logic Simulation (UG900), version 2022.2. — URL: https://docs.amd.com/r/2022.2-English/ug900-vivado-logic-simulation/ (дата обращения: 03.10.2026).",
           "Report Source")
    before(source_anchor,
           "9. AMD. Vivado Design Suite User Guide: Synthesis (UG901), version 2022.2. — URL: https://docs.amd.com/r/2022.2-English/ug901-vivado-synthesis/ (дата обращения: 03.10.2026).",
           "Report Source")

    doc.add_paragraph("Приложение Б", style="Report Appendix")
    doc.add_paragraph("Ключевые фрагменты RTL-модели и тестбенча", style="Report Appendix Title")
    doc.add_paragraph(
        "Ниже приведены операции, определяющие совпадение RTL-модели с "
        "целочисленным эталоном.", style="Report Body")
    doc.add_paragraph("Листинг Б.1 — Знаковое накопление и бинаризация слоя A", style="Report Listing Caption")
    code = doc.add_paragraph(style="Report Code")
    code.add_run(
        "accumulator = $signed(b_a[h]);\n"
        "for (j = 0; j < 5; j = j + 1) begin\n"
        "    product = $signed(input_features[39 - 8*j -: 8])\n"
        "            * $signed(w_a[j*8 + h]);\n"
        "    accumulator = accumulator + {{16{product[15]}}, product};\n"
        "end\n"
        "hidden_bits[h] = (accumulator >= 32'sd0);"
    )
    doc.add_paragraph("Листинг Б.2 — XNOR, popcount и пороговое решение", style="Report Listing Caption")
    code = doc.add_paragraph(style="Report Code")
    code.add_run(
        "for (h = 0; h < 8; h = h + 1)\n"
        "    matching_bits = matching_bits + (hidden_bits[h] ~^ w_b[h]);\n"
        "score = ($signed({28'b0, matching_bits}) <<< 1)\n"
        "      - 32'sd8 + $signed(b_b[0]);\n"
        "prediction_next = (score >= 32'sd0);"
    )
    for run in doc.paragraphs[-1].runs:
        run.font.size = Pt(9)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
