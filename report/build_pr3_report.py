"""Add measured PR3 results to the supplied report's PR1–PR2 continuation."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_combined_report import (
    before,
    evidence_figure_before,
    replace_text,
    table_before,
    toc_before,
)
from read_vivado_results import parse

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Отчет_ПР1_ПР2_Рябец_Головин_Гришин.docx"
OUTPUT = ROOT / "Отчет_ПР1_ПР2_ПР3_Рябец_Головин_Гришин.docx"
RESULTS = ROOT / "evidence" / "pr23"
REPO = ROOT.parent


def n(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def used(case: dict, resource: str) -> int:
    return case["routed"][resource]["used"]


def source_without_comments(filename: str) -> str:
    lines = (REPO / filename).read_text(encoding="utf-8").splitlines()
    return "\n".join(
        line.split("//", 1)[0].rstrip()
        for line in lines if not line.lstrip().startswith(("//", "#"))
    ).strip()


def first(doc, predicate):
    return next(paragraph for paragraph in doc.paragraphs if predicate(paragraph))


def pipeline_diagram_before(doc, anchor) -> None:
    labels = (
        "Входной регистр\n5 признаков",
        "Слой A\n5 × 8 MAC",
        "Регистр\n8 бит + valid",
        "Слой B\nXNOR + popcount",
        "Регистр\nответ + valid",
    )
    table = doc.add_table(rows=1, cols=len(labels))
    table.autofit = False
    for cell, label in zip(table.rows[0].cells, labels):
        cell.width = Cm(3.1)
        cell.text = label
        cell.paragraphs[0].style = doc.styles["Report Table Text"]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    anchor._p.addprevious(table._tbl)
    caption = before(anchor, "Рисунок 9.2 — Граница регистров между слоями сети", "Report Body")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(12)


def main() -> None:
    evidence = parse(RESULTS)
    c = evidence["cases"]
    baseline = c["baseline_20"]
    pipeline = c["pipeline_20"]
    proto = c["pr3_prototype_20"]
    proto_pipe = c["pr3_pipeline_20"]
    comparison = (
        f"При периоде 20 нс запас по установке вырос с "
        f"{n(baseline['wns_ns'])} до {n(pipeline['wns_ns'])} нс "
        f"(на {n(pipeline['wns_ns'] - baseline['wns_ns'])} нс). "
        f"При периоде 10 нс оба варианта выдержали ограничение: "
        f"WNS составил {n(c['baseline_10']['wns_ns'])} и "
        f"{n(c['pipeline_10']['wns_ns'])} нс соответственно. "
    )
    doc = Document(SOURCE)
    replace_text(doc.paragraphs[8], "ОТЧЁТ ПО ПРАКТИЧЕСКИМ РАБОТАМ № 1–3")
    replace_text(doc.paragraphs[11], "Тема: Программная модель, RTL-модель и прототип классификатора падений на ПЛИС")
    replace_text(
        first(doc, lambda p: p.text.startswith("Цель практических работ")),
        "Цель работы — разработать программный эталон классификатора падений, "
        "реализовать его на Verilog, проверить совпадение результатов и "
        "получить прототип для ПЛИС. После функциональной проверки "
        "исследуется влияние регистра между слоями на временные характеристики "
        "и использование ресурсов.",
    )
    intro_end = first(doc, lambda p: p.style.name == "Report Heading 1" and p.text.startswith("1 ПОСТАНОВКА ЗАДАЧИ"))
    before(intro_end, "– проверить верхний модуль с контрольными данными и выполнить размещение и трассировку;", "Report Body")
    before(intro_end, "– сравнить исходный и конвейерный вычислительные тракты при одинаковых входах и ограничениях.", "Report Body")

    toc_conclusion = first(doc, lambda p: p.style.name == "Report Contents" and p.text.startswith("ЗАКЛЮЧЕНИЕ"))
    toc_before(toc_conclusion, "9 ПРОТОТИП И КОНВЕЙЕРНАЯ АРХИТЕКТУРА\t14", toc_conclusion)
    toc_before(toc_conclusion, "10 ПРОВЕРКА И РЕЗУЛЬТАТЫ РЕАЛИЗАЦИИ\t16", toc_conclusion)
    replace_text(toc_conclusion, "ЗАКЛЮЧЕНИЕ\t19")
    replace_text(first(doc, lambda p: p.style.name == "Report Contents" and p.text.startswith("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")), "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ\t20")
    replace_text(first(doc, lambda p: p.style.name == "Report Contents" and p.text.startswith("ПРИЛОЖЕНИЕ А")), "ПРИЛОЖЕНИЕ А КЛЮЧЕВЫЕ ФРАГМЕНТЫ ПРОГРАММЫ\t21")
    toc_appendix_b = first(doc, lambda p: p.style.name == "Report Contents" and p.text.startswith("ПРИЛОЖЕНИЕ Б"))
    replace_text(toc_appendix_b, "ПРИЛОЖЕНИЕ Б ИСХОДНЫЙ КОД RTL-МОДЕЛИ\t26")
    toc_intro_end = first(doc, lambda p: p.style.name == "Report Heading 1" and p.text.strip() == "ВВЕДЕНИЕ")
    toc_before(toc_intro_end, "ПРИЛОЖЕНИЕ В КОД ПРОТОТИПА И КОНВЕЙЕРА\t29", toc_appendix_b)

    anchor = first(doc, lambda p: p.style.name == "Report Heading 1" and p.text.strip() == "ЗАКЛЮЧЕНИЕ")
    before(anchor, "9 ПРОТОТИП И КОНВЕЙЕРНАЯ АРХИТЕКТУРА", "Report Heading 1")
    before(anchor, "9.1 Верхний модуль и контрольные данные", "Report Heading 2")
    before(
        anchor,
        "Верхний модуль fpga_top хранит 64 входных слова и 64 ожидаемых бита, "
        "подготовленных программной моделью. Автомат IDLE–DRIVE–CHECK "
        "подаёт по одному вектору, сверяет ответ классификатора и ведёт "
        "счётчики проверок и расхождений. После последней проверки "
        "устанавливаются done и pass. Повторный импульс start запускает "
        "проверку заново. Для конвейерного классификатора реализован "
        "fpga_top_pipelined: он подаёт новый вектор каждый такт и "
        "задерживает индекс эталона вместе с ответом.",
        "Report Body",
    )
    before(anchor, "Таблица 9.1 — Сигналы верхнего модуля", "Report Table Caption")
    table_before(
        doc, anchor,
        ["Сигнал", "Разрядность", "Назначение"],
        [
            ["clk, rst_n", "1 + 1", "Тактирование и синхронный сброс"],
            ["start", "1", "Начало контрольного прогона"],
            ["busy, done, pass", "1 + 1 + 1", "Ход, завершение и результат"],
            ["checked_count, error_count", "7 + 7", "Проверенные векторы и расхождения"],
            ["last_prediction", "1", "Последний ответ сети"],
        ],
        [5.0, 2.8, 7.7],
    )
    before(
        anchor,
        "Контрольный набор подтверждает работу автомата, но он постоянен "
        "для синтезатора. Поэтому его схема не используется для оценки "
        "выигрыша в частоте: Vivado вправе упростить вычисления под 64 "
        "известных слова. Размещение верхнего модуля показано на рисунке 9.1.",
        "Report Body",
    )
    evidence_figure_before(anchor, "vivado_pr23_prototype_device.jpg", "Рисунок 9.1 — Размещение прототипа на кристалле")

    before(anchor, "9.2 Регистры между вычислительными слоями", "Report Heading 2")
    before(
        anchor,
        "В исходном классификаторе оба слоя образуют одну комбинационную "
        "цепь до регистра ответа. Для измерения этот тракт собран из тех же "
        "модулей bnn_layer_a и bnn_layer_b, что и конвейерный вариант. "
        "Восемь бит после слоя A и сигнал valid "
        "записываются в регистры; слой B формирует результат для следующего "
        "фронта. Регистр поставлен на границе слоёв, а не после каждого "
        "арифметического действия. Оба ядра принимают один вектор каждый "
        "такт, но конвейерный ответ появляется на один такт позже. "
        "Объявление reg в комбинационном always само по себе не означает "
        "физический регистр; его создаёт запись на posedge clk.",
        "Report Body",
    )
    pipeline_diagram_before(doc, anchor)

    before(anchor, "9.3 Условия сравнения", "Report Heading 2")
    before(
        anchor,
        "Оба варианта собраны в одном проекте Vivado 2022.2 для "
        "xc7a100tcsg324-1. Измерительные верхние модули имеют одинаковый "
        "40-битный вход, входной регистр, интерфейс valid и одни и те же "
        "параметры сети. Входы остаются переменными. Для каждого варианта "
        "выполнены синтез, размещение и трассировка при периодах 20 и 10 нс. "
        "Тактовый период ограничен соответствующим значением. "
        "Таким образом, сравнивается именно граница регистров "
        "вычислительного тракта. Выводы о таймингах относятся только к "
        "ограниченным внутренним синхронным путям.",
        "Report Body",
    )

    before(anchor, "10 ПРОВЕРКА И РЕЗУЛЬТАТЫ РЕАЛИЗАЦИИ", "Report Heading 1")
    before(anchor, "10.1 Функциональная проверка", "Report Heading 2")
    before(
        anchor,
        "Поведенческая симуляция исходного и конвейерного ядер сверила "
        "2331 вход с программным эталоном. Для обоих измерительных "
        "верхних модулей те же 2331 ответа проверены также в "
        "функциональной симуляции после синтеза. Самопроверяющиеся "
        "прототипы прошли контрольный набор из 64 слов; тестбенчи "
        "подтвердили обнаружение намеренно изменённого эталонного бита "
        "и повторный успешный запуск. При неизменённом эталоне "
        "расхождений нет.",
        "Report Body",
    )
    evidence_figure_before(anchor, "vivado_pr23_pipeline_pass.jpg", "Рисунок 10.1 — Сверка конвейерного классификатора с 2331 эталонным ответом в XSim")

    before(anchor, "10.2 Ресурсы и временные характеристики", "Report Heading 2")
    before(anchor, "Таблица 10.1 — Сравнение вычислительных трактов после трассировки", "Report Table Caption")
    table_before(
        doc, anchor,
        ["Вариант", "Период, нс", "Slice LUT", "Регистры", "WNS, нс"],
        [
            [label, n(c[key]["period_ns"]), str(used(c[key], "Slice LUTs")),
             str(used(c[key], "Slice Registers")), n(c[key]["wns_ns"])]
            for key, label in (
                ("baseline_20", "Исходный"),
                ("pipeline_20", "С регистром"),
                ("baseline_10", "Исходный"),
                ("pipeline_10", "С регистром"),
            )
        ],
        [3.8, 2.8, 2.8, 3.1, 3.0],
    )
    before(
        anchor,
        comparison +
        "При одинаковом периоде оба ядра принимают один вектор каждый такт. "
        "Промежуточный регистр добавил девять триггеров и один такт задержки "
        "отдельного ответа. Преимущество конвейера здесь — запас по таймингу, "
        "а не увеличение числа ответов за такт. Оба варианта прошли 10 нс "
        "(100 МГц) на ограниченных внутренних путях; точная предельная "
        "частота по этим двум прогонам не определена.",
        "Report Body",
    )
    prototype_caption = before(anchor, "Таблица 10.2 — Результаты прототипов с контрольной памятью", "Report Table Caption")
    prototype_caption.paragraph_format.page_break_before = True
    table_before(
        doc, anchor,
        ["Верхний модуль", "LUT", "Регистры", "Проложено цепей", "WNS при 20 нс"],
        [
            ["fpga_top", str(used(proto, "Slice LUTs")), str(used(proto, "Slice Registers")),
             f"{proto['routed_nets']}/{proto['routable_nets']}", n(proto["wns_ns"])],
            ["fpga_top_pipelined", str(used(proto_pipe, "Slice LUTs")),
             str(used(proto_pipe, "Slice Registers")),
             f"{proto_pipe['routed_nets']}/{proto_pipe['routable_nets']}", n(proto_pipe["wns_ns"])],
        ],
        [5.0, 1.6, 2.5, 3.4, 3.0],
    )
    before(
        anchor,
        "Оба прототипа полностью трассированы без ошибок маршрутизации. "
        "Эти значения включают автомат и постоянную контрольную память. "
        "Они характеризуют верхние модули, но не служат основанием "
        "для вывода о скорости вычислительного ядра.",
        "Report Body",
    )

    before(anchor, "10.3 Границы проверки", "Report Heading 2")
    rules = ", ".join(sorted(set(proto_pipe["drc_rules"])))
    before(
        anchor,
        f"Для прототипа не заданы внешние задержки у "
        f"{proto_pipe['input_ports_without_delay']} входов и "
        f"{proto_pipe['output_ports_without_delay']} выходов. "
        f"Отчёт DRC содержит {rules}. Плата не выбрана, поэтому "
        "не задавались физические выводы, стандарты ввода-вывода и "
        "напряжение конфигурации. Битовый поток не создавался и "
        "испытание на плате не проводилось. Подтверждены симуляция, "
        "синтез и физическая реализация для указанного кристалла.",
        "Report Body",
    )
    before(
        anchor,
        "В отдельном монолитном классификаторе Vivado предупредил об "
        "удалении одного промежуточного бита при оптимизации. Для сравнения "
        "вариантов использованы модульные вычислительные тракты; их ответы "
        "проверены на 2331 входе также после синтеза. Все четыре файла "
        "параметров сети успешно прочитаны при синтезе.",
        "Report Body",
    )

    conclusion = first(doc, lambda p: p.text.startswith("В практических работах № 1–2 создана"))
    replace_text(
        conclusion,
        "Построены программная модель, целочисленный RTL-классификатор и "
        "прототип на Artix-7. На 2331 эталонном векторе подтверждено "
        "совпадение вычислений с эталоном в поведенческой симуляции; "
        "для модульных вариантов — также после синтеза. Верхний модуль "
        "проверил 64 контрольных "
        "входа, обнаружение ошибки и повторный запуск. Размещение и "
        "трассировка в Vivado 2022.2 завершены для исходного и "
        "конвейерного вариантов. " + comparison +
        "Оценка относится к внутренним синхронным путям; "
        "назначение выводов и проверка на плате не выполнялись.",
    )
    source_anchor = first(doc, lambda p: p.style.name == "Report Appendix" and p.text.strip() == "Приложение А")
    before(
        source_anchor,
        "10. AMD. Vivado Design Suite User Guide: Implementation (UG904), "
        "version 2022.2. — URL: https://docs.amd.com/r/2022.2-English/ug904-vivado-implementation/ "
        "(дата обращения: 28.10.2026).",
        "Report Source",
    )

    doc.add_paragraph("Приложение В", style="Report Appendix")
    doc.add_paragraph("Код прототипа и конвейерного варианта", style="Report Appendix Title")
    doc.add_paragraph(
        "Приведены полные тексты вычислительных и верхних модулей, "
        "использованных для реализации и сравнения. Комментарии опущены.",
        style="Report Body",
    )
    sources = (
        ("В.1", "rtl/bnn_layer_a.v", "Слой A"),
        ("В.2", "rtl/bnn_layer_b.v", "Слой B"),
        ("В.3", "rtl/bnn_classifier_modular.v", "Классификатор без промежуточного регистра"),
        ("В.4", "rtl/bnn_classifier_pipelined.v", "Конвейерный классификатор"),
        ("В.5", "fpga/fpga_top.v", "Прототип с контрольными данными"),
        ("В.6", "fpga/fpga_top_pipelined.v", "Конвейерный прототип"),
        ("В.7", "fpga/bnn_benchmark.v", "Верхние модули для сравнения"),
    )
    for number, filename, title in sources:
        doc.add_paragraph(f"Листинг {number} — {title}", style="Report Listing Caption")
        paragraph = doc.add_paragraph(style="Report Code")
        paragraph.add_run(source_without_comments(filename)).font.size = Pt(9)

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Report Heading 1" and paragraph.text[:1].isdigit():
            paragraph.paragraph_format.page_break_before = False
            paragraph.paragraph_format.keep_with_next = True

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
